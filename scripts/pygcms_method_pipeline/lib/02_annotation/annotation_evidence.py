#!/usr/bin/env python3
# =============================================================================
# Purpose: Adjudicate spectral, RI, diagnostic-ion, and reference evidence.
# Associated manuscript:
# "A reproducible workflow for high-throughput pyrolysis gas
# chromatography–mass spectrometry analysis of soil organic matter".
# Usage, inputs, outputs, and parameters: see README.md.
# Responsibility: Compute candidate evidence and conservative feature-level decisions.
# =============================================================================

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import statistics
import time
from datetime import date
from html.parser import HTMLParser
from pathlib import Path
from typing import Any, Iterable
from urllib.parse import quote
from urllib.request import Request, urlopen

from openpyxl import load_workbook


def _feature_rt_for_annotation(feature):
    """Use the RT used for RI calculation, retaining support for older inputs."""
    for column in ("reference_corrected_rt_min", "corrected_tmean", "corrected_RT_min"):
        value = feature.get(column)
        if value is not None and str(value).strip():
            return value
    return ""


def _number(value: Any) -> float | None:
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _normalized_name(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "", _text(value).lower())


def _normalized_cas(value: Any) -> str:
    return re.sub(r"\D", "", _text(value))


def formula_identity_status(candidate_formula: Any, resolved_formula: Any) -> str:
    candidate = re.sub(r"\s+", "", _text(candidate_formula))
    resolved = re.sub(r"\s+", "", _text(resolved_formula))
    if not candidate or not resolved:
        return "unavailable"
    if candidate == resolved:
        return "matched"
    if re.search(r"(?:^|\d)D\d*", candidate) or re.search(r"(?:^|\d)D\d*", resolved):
        return "isotope_conflict"
    return "conflict"


def _valid_cas(value: Any) -> bool:
    match = re.fullmatch(r"(\d{2,7})-(\d{2})-(\d)", _text(value))
    if not match:
        return False
    digits = match.group(1) + match.group(2)
    checksum = sum((index + 1) * int(digit) for index, digit in enumerate(reversed(digits))) % 10
    return checksum == int(match.group(3))


def _empty_identity(method: str = "", query: str = "", status: str = "unresolved") -> dict[str, Any]:
    return {
        "identity_lookup_method": method,
        "identity_lookup_status": status,
        "identity_query": query,
        "resolved_pubchem_cid": "",
        "resolved_name": "",
        "resolved_CAS": "",
        "resolved_formula": "",
        "resolved_InChIKey": "",
        "identity_formula_status": "unavailable",
        "identity_source_url": "",
    }


def resolve_pubchem_identity(
    candidate: dict[str, Any],
    cache_dir: Path,
    online: bool,
    timeout: float = 20.0,
    delay: float = 0.12,
    opener: Any = urlopen,
) -> dict[str, Any]:
    queries: list[tuple[str, str]] = []
    inchikey = _text(candidate.get("InChIKey"))
    name = _text(candidate.get("Name"))
    if inchikey:
        queries.append(("inchikey", inchikey))
    if name:
        queries.append(("name", name))
    if not queries:
        return _empty_identity(status="no_query")

    # ``cache_dir`` is the final cache directory supplied by the pipeline.
    # Do not append another ``pubchem_identity`` component: doing so makes the
    # configured shared cache invisible to later runs.
    identity_cache = cache_dir
    identity_cache.mkdir(parents=True, exist_ok=True)
    last = _empty_identity(status="unresolved")
    for method, query in queries:
        digest = hashlib.sha256(f"{method}:{query.strip().lower()}".encode()).hexdigest()
        cache_path = identity_cache / f"{digest}.json"
        if cache_path.exists():
            cached = json.loads(cache_path.read_text(encoding="utf-8"))
            if cached.get("identity_lookup_status") == "resolved":
                return cached
            # A network exception is transient, not chemical evidence. Retry it
            # whenever online lookup is enabled instead of permanently poisoning
            # this identity query through the shared cache.
            if not (online and cached.get("identity_lookup_status") == "query_error"):
                last = cached
                continue
        if not online:
            last = _empty_identity(method, query, "not_cached_offline")
            continue

        encoded = quote(query, safe="")
        property_url = (
            f"https://pubchem.ncbi.nlm.nih.gov/rest/pug/compound/{method}/{encoded}/"
            "property/Title,MolecularFormula,InChIKey/JSON"
        )
        try:
            request = Request(property_url, headers={"User-Agent": "PyGCMS-RI-annotation/1.0 (research use)"})
            payload = json.loads(opener(request, timeout=timeout).read().decode("utf-8"))
            properties = payload.get("PropertyTable", {}).get("Properties", [])
            if len(properties) != 1:
                status = "not_found" if not properties else "ambiguous"
                last = _empty_identity(method, query, status)
                last["identity_source_url"] = property_url
            else:
                prop = properties[0]
                formula_status = formula_identity_status(candidate.get("Formula"), prop.get("MolecularFormula"))
                if formula_status != "matched":
                    last = _empty_identity(method, query, "rejected")
                    last.update({
                        "resolved_pubchem_cid": prop.get("CID", ""),
                        "resolved_name": _text(prop.get("Title")),
                        "resolved_formula": _text(prop.get("MolecularFormula")),
                        "resolved_InChIKey": _text(prop.get("InChIKey")),
                        "identity_formula_status": formula_status,
                        "identity_source_url": property_url,
                    })
                else:
                    cid = prop.get("CID", "")
                    synonym_url = f"https://pubchem.ncbi.nlm.nih.gov/rest/pug/compound/cid/{cid}/synonyms/JSON"
                    synonym_request = Request(synonym_url, headers={"User-Agent": "PyGCMS-RI-annotation/1.0 (research use)"})
                    synonym_payload = json.loads(opener(synonym_request, timeout=timeout).read().decode("utf-8"))
                    synonyms = synonym_payload.get("InformationList", {}).get("Information", [{}])[0].get("Synonym", [])
                    cas = next((_text(value) for value in synonyms if _valid_cas(value)), "")
                    last = _empty_identity(method, query, "resolved" if cas else "resolved_no_CAS")
                    last.update({
                        "resolved_pubchem_cid": cid,
                        "resolved_name": _text(prop.get("Title")),
                        "resolved_CAS": cas,
                        "resolved_formula": _text(prop.get("MolecularFormula")),
                        "resolved_InChIKey": _text(prop.get("InChIKey")),
                        "identity_formula_status": formula_status,
                        "identity_source_url": property_url,
                    })
        except Exception as exc:
            last = _empty_identity(method, query, "query_error")
            last["identity_source_url"] = property_url
            last["identity_error"] = type(exc).__name__
        last["identity_query_date"] = date.today().isoformat()
        cache_path.write_text(json.dumps(last, ensure_ascii=False, indent=2), encoding="utf-8")
        if delay > 0:
            time.sleep(delay)
        if last.get("identity_lookup_status") == "resolved":
            return last
    return last


def parse_reference_ions(value: Any) -> list[int]:
    ions: list[int] = []
    for token in re.findall(r"(?<![\d.])(\d+(?:\.\d+)?)(?![\d.])", _text(value)):
        number = float(token)
        if number.is_integer() and int(number) not in ions:
            ions.append(int(number))
    return ions


def parse_top_ions(value: Any, limit: int = 5) -> list[int]:
    parsed: list[tuple[int, float]] = []
    for mz, intensity in re.findall(r"(\d+(?:\.\d+)?)\s*[:;,]\s*(\d+(?:\.\d+)?)", _text(value)):
        mz_number = float(mz)
        if mz_number.is_integer():
            parsed.append((int(mz_number), float(intensity)))
    parsed.sort(key=lambda item: -item[1])
    return [mz for mz, _ in parsed[:limit]]


def top5_diagnostic_ions(
    defined_diagnostic_ions: Iterable[int],
    observed_top5: Iterable[int],
) -> list[int]:
    defined = {int(mz) for mz in defined_diagnostic_ions}
    measured_top5 = {int(mz) for mz in list(observed_top5)[:5]}
    return sorted(defined & measured_top5)


def diagnostic_ion_match_status(
    reference_ions: Iterable[int],
    observed_ions: Iterable[int],
) -> str:
    expected = sorted({int(mz) for mz in reference_ions})
    matched = sorted({int(mz) for mz in observed_ions})
    if not expected:
        return "Not defined"
    if not matched:
        return f"0/{len(expected)} — not observed in Top5"
    return f"{len(matched)}/{len(expected)} matched"


def parse_observed_ions(value: Any, limit: int = 5) -> list[int]:
    text = _text(value)
    if ":" in text:
        return parse_top_ions(text, limit=limit)
    return parse_reference_ions(text)[:limit]


def build_family_index(rows: Iterable[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    index: dict[str, dict[str, Any]] = {}
    for row in rows:
        class_code = _text(row.get("Class"))
        if class_code:
            references = _text(row.get("Reference(s)"))
            index[class_code] = {
                "family_name": _text(row.get("Class / family")),
                "diagnostic_ions": parse_reference_ions(row.get("Diagnostic ions (m/z)")),
                "literature_references": references,
                "literature_supported": bool(references),
            }
    return index


def is_db5_type(phase: Any) -> bool:
    value = re.sub(r"\s+", "", _text(phase).lower())
    if not value:
        return False
    explicit = (
        "db-5", "db5", "hp-5", "hp5", "rtx-5", "rtx5", "vf-5", "vf5",
        "spb-5", "spb5", "zb-5", "zb5", "bpx-5", "bpx5", "mdn-5", "mdn5",
        "cp-sil8", "cpsil8", "se-54", "se54", "ua-5", "ua5",
    )
    composition = bool(re.search(r"5%?(?:di)?phenyl", value))
    return any(marker in value for marker in explicit) or composition


def build_reference_index(rows: Iterable[dict[str, Any]]) -> dict[str, dict[str, list[dict[str, Any]]]]:
    by_cas: dict[str, list[dict[str, Any]]] = {}
    by_name: dict[str, list[dict[str, Any]]] = {}
    for original in rows:
        row = dict(original)
        cas = _normalized_cas(row.get("CAS"))
        if cas:
            by_cas.setdefault(cas, []).append(row)
        names = [_text(row.get("Compound"))]
        names.extend(re.split(r"\s*(?:;|\||/)\s*", _text(row.get("Synonym"))))
        for name in names:
            key = _normalized_name(name)
            if key:
                by_name.setdefault(key, []).append(row)
    return {"by_cas": by_cas, "by_name": by_name}


def _reference_payload(row: dict[str, Any], method: str) -> dict[str, Any]:
    phase = _text(row.get("RI_column") or row.get("RI Phase"))
    reference_ri = _number(row.get("RI_nonpolar (NIST)") or row.get("Non-polar RI")) if is_db5_type(phase) else None
    papers = int(_number(row.get("# papers") or row.get("Number of Papers")) or 0)
    soil_support = _text(row.get("Soil-matrix support") or row.get("Soil Support / Notes"))
    literature_supported = papers > 0 and not soil_support.lower().startswith("off-matrix")
    explicit_level = _text(row.get("Literature Support Level"))
    compound_name = _text(row.get("Compound"))
    if explicit_level:
        literature_level = explicit_level
    elif literature_supported and re.search(r"\bc\d+\s+(?:n-)?(?:alkane|alkene|fatty acid)\b", compound_name, re.I):
        literature_level = "Homologous Series"
    elif literature_supported:
        literature_level = "Compound"
    else:
        literature_level = "Unavailable"
    return {
        "reference_RI": reference_ri,
        "reference_RI_source": "SOM_reference_DB5" if reference_ri is not None else "",
        "reference_RI_columns": phase,
        "reference_match_method": method,
        "reference_compound": _text(row.get("Compound")),
        "reference_CAS": _text(row.get("CAS")),
        "reference_formula": _text(row.get("Formula")),
        "reference_class": _text(row.get("Class")),
        "reference_base_peak": int(_number(row.get("Base peak m/z") or row.get("Base Peak m/z"))) if _number(row.get("Base peak m/z") or row.get("Base Peak m/z")) is not None else None,
        "reference_diagnostic_ions": parse_reference_ions(row.get("Diagnostic ions m/z") or row.get("Diagnostic Ions m/z")),
        "reference_M_plus": int(_number(row.get("M+ (molecular ion)") or row.get("Molecular Ion (M+)"))) if _number(row.get("M+ (molecular ion)") or row.get("Molecular Ion (M+)")) is not None else None,
        "reference_n_papers": papers,
        "soil_matrix_support": soil_support,
        "literature_supported": literature_supported,
        "compound_literature_supported": literature_supported,
        "literature_support_level": literature_level,
    }


def match_reference_candidate(
    candidate: dict[str, Any],
    index: dict[str, dict[str, list[dict[str, Any]]]],
) -> dict[str, Any]:
    cas = _normalized_cas(candidate.get("CAS"))
    if cas and index["by_cas"].get(cas):
        return _reference_payload(index["by_cas"][cas][0], "CAS")
    name = _normalized_name(candidate.get("Name"))
    formula = _text(candidate.get("Formula"))
    matches = index["by_name"].get(name, [])
    if formula:
        formula_matches = [row for row in matches if _text(row.get("Formula")) == formula]
        if formula_matches:
            return _reference_payload(formula_matches[0], "name+formula")
    if matches:
        return _reference_payload(matches[0], "name")
    return {
        "reference_RI": None,
        "reference_RI_source": "",
        "reference_RI_columns": "",
        "reference_match_method": "",
        "reference_compound": "",
        "reference_CAS": "",
        "reference_formula": "",
        "reference_class": "",
        "reference_base_peak": None,
        "reference_diagnostic_ions": [],
        "reference_M_plus": None,
        "reference_n_papers": 0,
        "soil_matrix_support": "",
        "literature_supported": False,
        "compound_literature_supported": False,
        "family_literature_supported": False,
        "family_literature_references": "",
        "literature_support_level": "Unavailable",
    }


class _NISTTableParser(HTMLParser):
    """Collect each h3 heading and the first table that follows it."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.current_heading = ""
        self.pending_heading = ""
        self.in_heading = False
        self.in_table = False
        self.in_cell = False
        self.cell_kind = ""
        self.text_parts: list[str] = []
        self.current_row: list[tuple[str, str]] = []
        self.current_rows: list[list[tuple[str, str]]] = []
        self.tables: list[tuple[str, list[list[tuple[str, str]]]]] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        tag = tag.lower()
        if tag == "h3":
            self.in_heading = True
            self.text_parts = []
        elif tag == "table" and self.pending_heading:
            self.in_table = True
            self.current_rows = []
        elif self.in_table and tag in {"th", "td"}:
            self.in_cell = True
            self.cell_kind = tag
            self.text_parts = []

    def handle_data(self, data: str) -> None:
        if self.in_heading or self.in_cell:
            self.text_parts.append(data)

    def handle_endtag(self, tag: str) -> None:
        tag = tag.lower()
        if tag == "h3" and self.in_heading:
            self.current_heading = " ".join("".join(self.text_parts).split())
            self.pending_heading = self.current_heading
            self.in_heading = False
        elif self.in_table and tag in {"th", "td"} and self.in_cell:
            value = " ".join("".join(self.text_parts).split())
            self.current_row.append((self.cell_kind, value))
            self.in_cell = False
        elif self.in_table and tag == "tr":
            if self.current_row:
                self.current_rows.append(self.current_row)
            self.current_row = []
        elif tag == "table" and self.in_table:
            self.tables.append((self.pending_heading, self.current_rows))
            self.pending_heading = ""
            self.in_table = False


def parse_nist_db5_ri_html(html: str) -> list[dict[str, Any]]:
    parser = _NISTTableParser()
    parser.feed(html)
    records: list[dict[str, Any]] = []
    for label, rows in parser.tables:
        label_lower = label.lower()
        if "ri" not in label_lower or "lee" in label_lower or "isothermal" in label_lower:
            continue
        if "temperature ramp" not in label_lower and "temperature program" not in label_lower:
            continue
        headers = [value for kind, value in rows[0] if kind == "th"] if rows else []
        for row in rows:
            cells = [value for kind, value in row if kind == "td"]
            if len(cells) < 3:
                continue
            phase = cells[1]
            value = _number(cells[2])
            if value is None or not is_db5_type(phase):
                continue
            record = {"RI_type": label, "column_type": cells[0], "stationary_phase": phase, "RI": value}
            for idx, header in enumerate(headers[3:], start=3):
                if idx < len(cells):
                    record[header] = cells[idx]
            records.append(record)
    return records


def _read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        return list(csv.DictReader(handle))


def _write_csv(path: Path, rows: list[dict[str, Any]], preferred_fields: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fields = list(preferred_fields)
    for row in rows:
        for key in row:
            if key not in fields:
                fields.append(key)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def _sheet_rows(sheet: Any, header_row: int = 1) -> list[dict[str, Any]]:
    iterator = sheet.iter_rows(values_only=True)
    for _ in range(header_row - 1):
        next(iterator)
    headers = [_text(value) for value in next(iterator)]
    return [dict(zip(headers, values)) for values in iterator if any(value not in (None, "") for value in values)]


def _pipeline_library_rows(workbook: Any) -> list[dict[str, Any]]:
    pipeline_rows = _sheet_rows(workbook["Pipeline_ion_library"])
    master_rows = (
        _sheet_rows(workbook["Compound_master"])
        if "Compound_master" in workbook.sheetnames
        else []
    )
    alias_rows = (
        _sheet_rows(workbook["Compound_alias_map"])
        if "Compound_alias_map" in workbook.sheetnames
        else []
    )
    master_by_id = {
        _text(row.get("Compound_ID")): row
        for row in master_rows
        if _text(row.get("Compound_ID"))
    }
    aliases_by_id: dict[str, list[str]] = {}
    for compound_id, master in master_by_id.items():
        aliases_by_id[compound_id] = [
            part.strip()
            for part in _text(master.get("Reported_names")).split(";")
            if part.strip()
        ]
    for alias in alias_rows:
        compound_id = _text(alias.get("Compound_ID"))
        reported_name = _text(alias.get("Reported_name"))
        if compound_id and reported_name:
            aliases_by_id.setdefault(compound_id, []).append(reported_name)

    compounds: list[dict[str, Any]] = []
    for pipeline_row in pipeline_rows:
        compound_id = _text(pipeline_row.get("Compound_ID"))
        master = master_by_id.get(compound_id, {})
        canonical_name = _text(
            pipeline_row.get("Canonical_name") or master.get("Canonical_name")
        )
        seen_aliases: set[str] = set()
        synonyms: list[str] = []
        canonical_key = _normalized_name(canonical_name)
        for alias in aliases_by_id.get(compound_id, []):
            key = _normalized_name(alias)
            if not key or key == canonical_key or key in seen_aliases:
                continue
            seen_aliases.add(key)
            synonyms.append(alias)
        compounds.append({
            "Compound_ID": compound_id,
            "Compound": canonical_name,
            "Synonym": "; ".join(synonyms),
            "Class": _text(
                pipeline_row.get("Compound_class")
                or master.get("Compound_classes")
            ),
            "Base peak m/z": pipeline_row.get("Base_mz"),
            "Diagnostic ions m/z": _text(
                pipeline_row.get("Diagnostic_ions_mz")
            ),
            "# papers": int(_number(master.get("Reference_count")) or 0),
            "References": _text(master.get("References")),
        })
    return compounds


def load_reference_workbook(path: Path) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    compound_sheet = next(
        (name for name in ("Compounds library", "Identification library") if name in workbook.sheetnames),
        None,
    )
    if compound_sheet is not None:
        compounds = _sheet_rows(workbook[compound_sheet])
    elif "Pipeline_ion_library" in workbook.sheetnames:
        compounds = _pipeline_library_rows(workbook)
    else:
        raise ValueError(
            "Missing 'Pipeline_ion_library', 'Compounds library', or "
            f"'Identification library' sheet in {path}"
        )
    families: list[dict[str, Any]] = []
    if "Family diagnostic ions" in workbook.sheetnames:
        family_sheet = workbook["Family diagnostic ions"]
        first_row = [_text(value) for value in next(family_sheet.iter_rows(values_only=True))]
        header_row = 1 if "Class" in first_row else 4
        families = _sheet_rows(family_sheet, header_row=header_row)
    return compounds, families


def load_reference_rows(path: Path) -> list[dict[str, Any]]:
    return load_reference_workbook(path)[0]


def load_alkane_ladder(path: Path) -> list[tuple[int, float]]:
    rows = _read_csv(path)
    points = []
    for row in rows:
        carbon = _number(row.get("carbon_number"))
        rt = _number(row.get("RT_min"))
        if carbon is not None and rt is not None:
            points.append((int(carbon), rt))
    return sorted(points)


def nist_webbook_url(cas: Any) -> str:
    digits = _normalized_cas(cas)
    return f"https://webbook.nist.gov/cgi/cbook.cgi?ID=C{digits}&Mask=2200" if digits else ""


def get_nist_db5_records(
    cas: Any,
    cache_dir: Path,
    online: bool,
    timeout: float = 20.0,
    delay: float = 0.12,
) -> dict[str, Any]:
    digits = _normalized_cas(cas)
    if not digits:
        return {"records": [], "source_url": "", "cache_status": "no_CAS"}
    cache_dir.mkdir(parents=True, exist_ok=True)
    cache_path = cache_dir / f"{digits}.json"
    if cache_path.exists():
        cached = json.loads(cache_path.read_text(encoding="utf-8"))
        if not (online and cached.get("cache_status") == "query_error"):
            return cached
    url = nist_webbook_url(cas)
    if not online:
        return {"records": [], "source_url": url, "cache_status": "not_cached_offline"}
    request = Request(url, headers={"User-Agent": "PyGCMS-RI-annotation/1.0 (research use)"})
    try:
        html = urlopen(request, timeout=timeout).read().decode("utf-8", errors="replace")
        records = parse_nist_db5_ri_html(html)
        payload = {
            "CAS": _text(cas),
            "source_url": url,
            "query_date": date.today().isoformat(),
            "cache_status": "downloaded",
            "records": records,
        }
    except Exception as exc:
        return {
            "CAS": _text(cas),
            "source_url": url,
            "query_date": date.today().isoformat(),
            "cache_status": "query_error",
            "error": type(exc).__name__,
            "records": [],
        }
    cache_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    if delay > 0:
        time.sleep(delay)
    return payload


def attach_reference_evidence(
    candidate: dict[str, Any],
    reference_index: dict[str, dict[str, list[dict[str, Any]]]],
    cache_dir: Path,
    online: bool,
    timeout: float,
    delay: float,
    pubchem_cache_dir: Path | None = None,
    pubchem_online: bool | None = None,
    pubchem_timeout: float | None = None,
    pubchem_delay: float | None = None,
    identity_resolver: Any = resolve_pubchem_identity,
) -> dict[str, Any]:
    row = dict(candidate)
    row.update(_empty_identity())
    local = match_reference_candidate(row, reference_index)
    effective_cas = _text(row.get("CAS"))
    if local.get("reference_RI") is None and not _normalized_cas(effective_cas):
        identity = identity_resolver(
            candidate,
            pubchem_cache_dir or cache_dir,
            online=online if pubchem_online is None else pubchem_online,
            timeout=timeout if pubchem_timeout is None else pubchem_timeout,
            delay=delay if pubchem_delay is None else pubchem_delay,
        )
        row.update(identity)
        if identity.get("identity_lookup_status") == "resolved" and _valid_cas(identity.get("resolved_CAS")):
            effective = dict(candidate)
            effective["CAS"] = identity["resolved_CAS"]
            effective["Name"] = identity.get("resolved_name") or candidate.get("Name")
            effective["Formula"] = identity.get("resolved_formula") or candidate.get("Formula")
            local = match_reference_candidate(effective, reference_index)
            effective_cas = identity["resolved_CAS"]
            if local.get("reference_match_method") == "CAS":
                local["reference_match_method"] = "resolved_CAS"
    row.update(local)
    embedded_ri = _number(candidate.get("embedded_SemiStdNP_RI"))
    if "nist" in _text(candidate.get("library")).lower() and embedded_ri is not None:
        embedded_values = [
            value for value in (
                _number(token) for token in re.split(r"[;,|]", _text(candidate.get("embedded_SemiStdNP_values")))
            ) if value is not None
        ]
        row["reference_RI"] = embedded_ri
        row["reference_RI_source"] = "NIST2020_MSP_SemiStdNP_median"
        row["reference_RI_columns"] = "SemiStdNP"
        row["reference_match_method"] = "NIST2020_InChIKey_or_CAS"
        row["n_reference_RI_records"] = int(_number(candidate.get("embedded_SemiStdNP_count")) or len(embedded_values) or 1)
        row["reference_RI_min"] = _number(candidate.get("embedded_SemiStdNP_min"))
        row["reference_RI_max"] = _number(candidate.get("embedded_SemiStdNP_max"))
        if row["reference_RI_min"] is None:
            row["reference_RI_min"] = min(embedded_values) if embedded_values else embedded_ri
        if row["reference_RI_max"] is None:
            row["reference_RI_max"] = max(embedded_values) if embedded_values else embedded_ri
        row["reference_RI_source_url"] = ""
        row["RI_cache_status"] = "not_queried_embedded_RI"
        return row
    if row.get("reference_RI") is not None:
        row["n_reference_RI_records"] = 1
        row["reference_RI_min"] = row["reference_RI"]
        row["reference_RI_max"] = row["reference_RI"]
        row["reference_RI_source_url"] = ""
        return row

    cas = effective_cas or row.get("reference_CAS")
    result = get_nist_db5_records(cas, cache_dir, online=online, timeout=timeout, delay=delay)
    values = [record["RI"] for record in result.get("records", []) if _number(record.get("RI")) is not None]
    if values:
        row["reference_RI"] = statistics.median(values)
        row["reference_RI_source"] = "NIST_WebBook_DB5_median"
        row["reference_RI_columns"] = "; ".join(sorted({_text(record.get("stationary_phase")) for record in result["records"]}))
    row["n_reference_RI_records"] = len(values)
    row["reference_RI_min"] = min(values) if values else ""
    row["reference_RI_max"] = max(values) if values else ""
    row["reference_RI_source_url"] = result.get("source_url", "")
    row["RI_cache_status"] = result.get("cache_status", "")
    return row


def run_identification(
    *,
    top_hits_path: Path,
    feature_metadata_path: Path,
    alkane_path: Path,
    reference_workbook_path: Path,
    candidate_output_path: Path,
    decision_output_path: Path,
    cache_dir: Path,
    pubchem_cache_dir: Path | None = None,
    online: bool = True,
    pubchem_online: bool | None = None,
    timeout: float = 20.0,
    delay: float = 0.12,
    pubchem_timeout: float = 20.0,
    pubchem_delay: float = 0.12,
    support_window: float = 20.0,
    weak_window: float = 50.0,
    spectral_auto_threshold: float = 850.0,
    decision_policy: str = "legacy_ion_first",
    primary_top_n: int = 5,
    rescue_top_n: int = 20,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    top_hits = _read_csv(top_hits_path)
    features = _read_csv(feature_metadata_path)
    ladder = load_alkane_ladder(alkane_path)
    reference_rows, family_rows = load_reference_workbook(reference_workbook_path)
    reference_index = build_reference_index(reference_rows)
    family_index = build_family_index(family_rows)
    hits_by_feature: dict[str, list[dict[str, Any]]] = {}
    for hit in top_hits:
        hits_by_feature.setdefault(_text(hit.get("GlobalFeatureID")), []).append(hit)

    candidate_rows: list[dict[str, Any]] = []
    decision_rows: list[dict[str, Any]] = []
    for feature in features:
        feature_id = _text(feature.get("GlobalFeatureID"))
        rt = _number(_feature_rt_for_annotation(feature))
        ri_calc = resolve_feature_ri(feature, ladder)
        observed_base_mz = _number(feature.get("BaseMz"))
        observed_top5 = parse_top_ions(feature.get("TopIons"), limit=5)

        enriched = [
            attach_ion_evidence(
                attach_reference_evidence(
                    hit,
                    reference_index,
                    cache_dir=cache_dir,
                    online=online,
                    timeout=timeout,
                    delay=delay,
                    pubchem_cache_dir=pubchem_cache_dir,
                    pubchem_online=pubchem_online,
                    pubchem_timeout=pubchem_timeout,
                    pubchem_delay=pubchem_delay,
                ),
                observed_base_mz=observed_base_mz,
                observed_top_ions=observed_top5,
                family_index=family_index,
            )
            for hit in hits_by_feature.get(feature_id, [])
        ]
        homolog = detect_homologous_series(
            measured_ri=ri_calc["measured_RI"],
            ri_calculation_status=ri_calc["RI_calculation_status"],
            observed_base_mz=observed_base_mz,
            observed_top_ions=observed_top5,
            candidates=enriched,
        )
        audited = audit_candidates(
            enriched,
            measured_ri=ri_calc["measured_RI"],
            support_window=support_window,
            weak_window=weak_window,
            spectral_auto_threshold=spectral_auto_threshold,
        )
        for row in audited:
            row.update({
                "GlobalFeatureID": feature_id,
                "corrected_RT_min": rt if rt is not None else "",
                "preceding_alkane": ri_calc["preceding_alkane"],
                "following_alkane": ri_calc["following_alkane"],
                "RI_calculation_status": ri_calc["RI_calculation_status"],
                "RI_confidence": ri_calc["RI_confidence"],
                "observed_base_mz": observed_base_mz if observed_base_mz is not None else "",
                **homolog,
            })
        candidate_rows.extend(audited)
        if decision_policy == "two_stage_top20_rescue":
            decision = select_two_stage_feature_decision(
                audited,
                ri_calculation_status=ri_calc["RI_calculation_status"],
                spectral_auto_threshold=spectral_auto_threshold,
                primary_top_n=primary_top_n,
                rescue_top_n=rescue_top_n,
            )
        else:
            decision = select_feature_decision(
                audited,
                ri_calculation_status=ri_calc["RI_calculation_status"],
                spectral_auto_threshold=spectral_auto_threshold,
                decision_policy=decision_policy,
            )
        if homolog["homologous_series_status"] != "not_detected":
            carbon = int(homolog["homologous_carbon_number"])
            if homolog["homologous_series_class"] == "n_alkane":
                formula = f"C{carbon}H{2 * carbon + 2}"
            elif homolog["homologous_series_class"] == "alkene":
                formula = f"C{carbon}H{2 * carbon}"
            else:
                formula = ""
            decision.update({
                "automated_status": homolog["homologous_series_status"],
                "decision_round": "homologous_series",
                "selected_candidate": homolog["homologous_series_assignment"],
                "selected_formula": formula,
                "selected_library": "homologous_series_rule",
                "automatic_identification_eligible": True,
                "selected_literature_support_level": (
                    "Homologous Series"
                    if family_index.get("Al", {}).get("literature_supported")
                    else "Unavailable"
                ),
                "selected_matched_diagnostic_ions": ";".join(
                    str(ion) for ion in matched_homologous_ions(
                        homolog["homologous_series_class"], observed_top5
                    )
                ),
                **homolog,
            })
        decision.update({
            "GlobalFeatureID": feature_id,
            "corrected_RT_min": rt if rt is not None else "",
            "measured_RI": ri_calc["measured_RI"] if ri_calc["measured_RI"] is not None else "",
            "preceding_alkane": ri_calc["preceding_alkane"],
            "following_alkane": ri_calc["following_alkane"],
            "RI_calculation_status": ri_calc["RI_calculation_status"],
            "RI_confidence": ri_calc["RI_confidence"],
        })
        decision_rows.append(decision)

    _write_csv(candidate_output_path, candidate_rows, [
        "GlobalFeatureID", "corrected_RT_min", "measured_RI", "preceding_alkane", "following_alkane",
        "RI_calculation_status", "RI_confidence", "library", "rank", "score_0_1000", "spectral_cosine",
        "matched_ion_count", "query_ion_count", "matched_ion_fraction",
        "matched_intensity_fraction", "base_peak_matched", "possible_coelution", "Name",
        "spectral_evidence_status", "automatic_identification_eligible",
        "InChIKey", "CAS", "Formula", "reference_RI", "delta_RI", "abs_delta_RI", "RI_status",
        "identity_lookup_method", "identity_lookup_status", "identity_query", "resolved_pubchem_cid",
        "resolved_name", "resolved_CAS", "resolved_formula", "resolved_InChIKey",
        "identity_formula_status", "identity_source_url",
        "reference_RI_source", "n_reference_RI_records", "reference_RI_min", "reference_RI_max",
        "embedded_SemiStdNP_RI", "embedded_SemiStdNP_values", "embedded_SemiStdNP_count",
        "embedded_SemiStdNP_min", "embedded_SemiStdNP_max", "nist_hit_record_count", "nist_source_files",
        "reference_RI_columns", "reference_match_method", "reference_compound", "reference_n_papers",
        "soil_matrix_support", "literature_supported", "compound_literature_supported",
        "family_literature_supported", "family_literature_references", "literature_support_level",
        "reference_RI_source_url", "observed_base_mz",
        "observed_top5_ions", "reference_diagnostic_ions_text",
        "observed_diagnostic_ions", "diagnostic_ion_match",
        "reference_base_peak", "base_peak_status",
        "expected_compound_diagnostic_ions", "expected_additional_diagnostic_ions",
        "matched_additional_diagnostic_ions", "n_additional_diagnostic_ions_matched", "compound_ion_status",
        "reference_class", "family_name", "expected_family_ions", "matched_family_ions",
        "n_family_ions_matched", "family_ion_status", "family_base_peak_status",
        "family_evidence_reason", "reference_M_plus", "M_plus_status",
        "homologous_series_class", "homologous_carbon_number", "homologous_RI_residual",
        "homologous_series_assignment", "homologous_series_status",
    ])
    _write_csv(decision_output_path, decision_rows, [
        "GlobalFeatureID", "corrected_RT_min", "measured_RI", "preceding_alkane", "following_alkane",
        "RI_calculation_status", "RI_confidence", "automated_status", "selected_candidate", "selected_formula", "selected_CAS",
        "selected_library", "selected_match_score", "selected_reference_RI", "selected_delta_RI",
        "selected_base_peak_status", "selected_compound_ion_status", "selected_matched_compound_ions",
        "selected_family", "selected_family_ion_status", "selected_matched_family_ions", "selected_M_plus_status",
        "selected_literature_support_level",
        "selected_matched_diagnostic_ions", "selected_reference_diagnostic_ions",
        "selected_observed_diagnostic_ions", "selected_diagnostic_ion_match",
        "selected_identification_evidence",
        "spectral_auto_threshold", "automatic_identification_eligible", "RI_ambiguity_gap",
        "decision_round", "primary_automated_status",
        "primary_selected_candidate", "rescue_attempted",
        "rescue_automated_status", "rescue_selected_candidate",
        "homologous_series_class", "homologous_carbon_number", "homologous_RI_residual",
        "homologous_series_assignment", "homologous_series_status",
        "n_candidates_retained", "n_auto_eligible_candidates", "n_RI_supported_candidates", "candidate_group",
    ])
    return candidate_rows, decision_rows


def calculate_measured_ri(retention_time: float, ladder: Iterable[tuple[int, float]]) -> dict[str, Any]:
    points = sorted((int(carbon), float(rt)) for carbon, rt in ladder)
    if len(points) < 2:
        raise ValueError("At least two n-alkane retention times are required")

    if retention_time < points[0][1]:
        return {
            "measured_RI": None,
            "preceding_alkane": "",
            "following_alkane": points[0][0],
            "RI_calculation_status": "outside_calibration_range",
        }
    if retention_time > points[-1][1]:
        return {
            "measured_RI": None,
            "preceding_alkane": points[-1][0],
            "following_alkane": "",
            "RI_calculation_status": "outside_calibration_range",
        }

    lower, upper = points[0], points[1]
    for left, right in zip(points, points[1:]):
        if left[1] <= retention_time <= right[1]:
            lower, upper = left, right
            break

    carbon, lower_rt = lower
    next_carbon, upper_rt = upper
    if next_carbon != carbon + 1:
        raise ValueError("The bracketing n-alkanes must have adjacent carbon numbers")
    if upper_rt <= lower_rt:
        raise ValueError("n-Alkane retention times must be strictly increasing")
    measured = 100.0 * (carbon + (retention_time - lower_rt) / (upper_rt - lower_rt))
    return {
        "measured_RI": measured,
        "preceding_alkane": carbon,
        "following_alkane": next_carbon,
        "RI_calculation_status": "interpolated",
    }


def resolve_feature_ri(
    feature: dict[str, Any], ladder: Iterable[tuple[int, float]]
) -> dict[str, Any]:
    if "measured_RI" in feature:
        measured = _number(feature.get("measured_RI"))
        status = _text(
            feature.get("RI_calculation_status") or feature.get("RI_status")
        )
        if not status:
            status = "interpolated" if measured is not None else "unavailable"
        return {
            "measured_RI": measured,
            "preceding_alkane": _number(
                feature.get("preceding_alkane")
                or feature.get("lower_alkane_carbon")
            ) or "",
            "following_alkane": _number(
                feature.get("following_alkane")
                or feature.get("upper_alkane_carbon")
            ) or "",
            "RI_calculation_status": status,
            "RI_confidence": _text(feature.get("RI_confidence"))
            or "precomputed_measured_ri",
        }

    retention_time = _number(_feature_rt_for_annotation(feature))
    if retention_time is None:
        return {
            "measured_RI": None,
            "preceding_alkane": "",
            "following_alkane": "",
            "RI_calculation_status": "RT_unavailable",
            "RI_confidence": "legacy_rt_ladder",
        }
    result = calculate_measured_ri(retention_time, ladder)
    result["RI_confidence"] = "legacy_rt_ladder"
    return result


def ri_status(delta_ri: float | None, support_window: float = 20.0, weak_window: float = 50.0) -> str:
    if delta_ri is None:
        return "unavailable"
    difference = abs(delta_ri)
    if difference <= support_window:
        return "supported"
    if difference <= weak_window:
        return "weak_support"
    return "mismatch"


def audit_candidates(
    candidates: Iterable[dict[str, Any]],
    measured_ri: float | None,
    support_window: float = 20.0,
    weak_window: float = 50.0,
    spectral_auto_threshold: float = 850.0,
) -> list[dict[str, Any]]:
    audited: list[dict[str, Any]] = []
    for candidate in candidates:
        row = dict(candidate)
        reference_ri = _number(row.get("reference_RI"))
        delta = None if measured_ri is None or reference_ri is None else measured_ri - reference_ri
        row["measured_RI"] = measured_ri
        row["reference_RI"] = reference_ri
        row["delta_RI"] = delta
        row["abs_delta_RI"] = None if delta is None else abs(delta)
        row["RI_status"] = ri_status(delta, support_window=support_window, weak_window=weak_window)
        eligible = _score(row) >= spectral_auto_threshold
        row["spectral_evidence_status"] = "strong" if eligible else "below_auto_threshold"
        row["automatic_identification_eligible"] = eligible
        audited.append(row)
    return audited


def _formula_carbon_hydrogen(formula: Any) -> tuple[int | None, int | None]:
    text = _text(formula)
    carbon = re.search(r"(?:^|[^A-Za-z])C(\d*)", text)
    hydrogen = re.search(r"H(\d*)", text)
    if not carbon:
        return None, None
    c = int(carbon.group(1) or "1")
    h = int(hydrogen.group(1) or "1") if hydrogen else None
    return c, h


def detect_homologous_series(
    *,
    measured_ri: float | None,
    ri_calculation_status: str,
    observed_base_mz: Any,
    observed_top_ions: Iterable[int],
    candidates: Iterable[dict[str, Any]],
    ri_window: float = 20.0,
) -> dict[str, Any]:
    result = {
        "homologous_series_class": "",
        "homologous_carbon_number": "",
        "homologous_RI_residual": "",
        "homologous_series_assignment": "",
        "homologous_series_status": "not_detected",
    }
    if measured_ri is None or ri_calculation_status not in {
        "interpolated", "ladder_anchor"
    }:
        return result
    top5 = set(list(observed_top_ions)[:5])
    base = _number(observed_base_mz)
    carbon_from_ri = int(round(measured_ri / 100.0))
    residual = measured_ri - 100.0 * carbon_from_ri

    if base is not None and int(base) == 57 and {57, 71, 85}.issubset(top5) and abs(residual) <= ri_window:
        return {
            "homologous_series_class": "n_alkane",
            "homologous_carbon_number": carbon_from_ri,
            "homologous_RI_residual": residual,
            "homologous_series_assignment": f"C{carbon_from_ri} n-alkane",
            "homologous_series_status": "homologous_series_n_alkane",
        }

    if ri_calculation_status != "interpolated":
        return result

    candidate_rows = list(candidates)
    alkene_supported = False
    for row in candidate_rows:
        name = _text(row.get("Name")).lower()
        class_code = _text(row.get("reference_class")).lower()
        c, h = _formula_carbon_hydrogen(row.get("Formula"))
        name_support = bool(
            "alkene" in name
            or "olefin" in name
            or re.search(r"(?:^|[^a-z])[a-z0-9,()\-]*(?:diene|triene|ene)\b", name)
        )
        class_support = class_code in {"alkene", "alkenes", "olefin", "olefins"}
        formula_compatible = c is None or h is None or h == 2 * c
        if (name_support or class_support) and formula_compatible:
            alkene_supported = True
            break
    if {55, 69, 83}.issubset(top5) and alkene_supported and abs(residual) <= ri_window:
        return {
            "homologous_series_class": "alkene",
            "homologous_carbon_number": carbon_from_ri,
            "homologous_RI_residual": residual,
            "homologous_series_assignment": f"C{carbon_from_ri} alkene",
            "homologous_series_status": "homologous_series_alkene",
        }

    acid_candidates = []
    for row in candidate_rows:
        name = _text(row.get("Name")).lower()
        class_code = _text(row.get("reference_class")).lower()
        if "acid" in name or class_code in {"fa", "fatty acid", "lipid"}:
            c, _ = _formula_carbon_hydrogen(row.get("Formula") or row.get("reference_formula"))
            if c is not None:
                acid_candidates.append((_score(row), c))
    if {60, 73}.issubset(top5) and acid_candidates:
        carbon = sorted(acid_candidates, reverse=True)[0][1]
        return {
            "homologous_series_class": "fatty_acid",
            "homologous_carbon_number": carbon,
            "homologous_RI_residual": "",
            "homologous_series_assignment": f"C{carbon} fatty acid",
            "homologous_series_status": "homologous_series_fatty_acid",
        }
    return result


def matched_homologous_ions(series_class: str, observed_top_ions: Iterable[int]) -> list[int]:
    expected = {
        "n_alkane": (57, 71, 85),
        "alkene": (55, 69, 83),
        "fatty_acid": (60, 73, 129),
    }.get(series_class, ())
    observed = set(list(observed_top_ions)[:5])
    return [ion for ion in expected if ion in observed]


def _join_ions(ions: Iterable[int]) -> str:
    return ";".join(str(ion) for ion in ions)


def _classify_family_ion_evidence(
    row: dict[str, Any],
    class_code: str,
    family_name: str,
    observed_base: float | None,
    top5: list[int],
    family_expected: list[int],
    family_matched: list[int],
    molecular_status: str,
) -> tuple[str, str, str]:
    if not class_code or not family_expected:
        return "unavailable", "unavailable", "family_reference_unavailable"

    base = int(observed_base) if observed_base is not None else None
    top3 = set(top5[:3])
    top5_set = set(top5)
    name = _text(row.get("Name")).lower()
    literature_or_structure = bool(
        row.get("literature_supported")
        or _text(row.get("reference_match_method"))
        or _text(row.get("reference_compound"))
    )

    if class_code == "Al":
        straight_alkane_names = (
            "heptane", "octane", "nonane", "decane", "undecane", "dodecane",
            "tridecane", "tetradecane", "pentadecane", "hexadecane", "heptadecane",
            "octadecane", "nonadecane", "eicosane", "heneicosane", "docosane",
            "tricosane", "tetracosane", "pentacosane", "hexacosane", "heptacosane",
            "octacosane", "nonacosane", "triacontane", "hentriacontane",
        )
        is_straight_alkane = name.strip(" .") in straight_alkane_names or "n-alkane" in name
        is_fatty_acid = " acid" in name or name.endswith("acid")
        is_alkene = not is_fatty_acid and ("alkene" in name or bool(re.search(r"\b\d*-[a-z]+ene\b", name)))

        if is_straight_alkane:
            supported = base == 57 and {71, 85}.issubset(top5_set)
            return (
                "supported" if supported else "not_supported",
                "matched" if base == 57 else "conflict",
                "n_alkane_base57_plus_71_85" if supported else "n_alkane_rule_failed",
            )
        if is_alkene:
            supported = base in {55, 69} and {55, 69, 83}.issubset(top5_set)
            return (
                "supported" if supported else "not_supported",
                "matched" if base in {55, 69} else "conflict",
                "alkene_base55_or69_plus_55_69_83" if supported else "alkene_rule_failed",
            )
        if is_fatty_acid:
            base_ok = 60 in top3
            fragment_ok = 73 in top5_set and (129 in top5_set or molecular_status == "matched")
            supported = base_ok and fragment_ok
            return (
                "supported" if supported else "not_supported",
                "matched" if base == 60 else "top3_compatible" if base_ok else "conflict",
                "fatty_acid_strong60_plus_73_and_129_or_Mplus" if supported else "fatty_acid_rule_failed",
            )
        return "not_supported", "unavailable", "aliphatic_subfamily_not_resolved"

    if class_code == "Ps":
        structural_name = any(token in name for token in ("furan", "furfural", "pyran", "maltol", "sugar"))
        supported = len(family_matched) >= 3 and (literature_or_structure or structural_name)
        base_status = "compatible" if base in set(family_expected) else "not_required"
        return (
            "supported" if supported else "not_supported",
            base_status,
            "carbohydrate_furan_three_markers" if supported else "carbohydrate_furan_rule_failed",
        )

    if class_code in {"Lg", "Ls"}:
        supported = len(family_matched) >= 3 and literature_or_structure
        return (
            "supported" if supported else "not_supported",
            "compatible" if base in set(family_expected) else "not_required",
            "lignin_three_markers_plus_structure" if supported else "lignin_rule_failed",
        )

    supported = len(family_matched) >= 3 and literature_or_structure and base in set(family_expected)
    return (
        "supported" if supported else "not_supported",
        "compatible" if base in set(family_expected) else "conflict",
        "generic_three_markers_plus_structure_and_base" if supported else "generic_family_rule_failed",
    )


def attach_ion_evidence(
    candidate: dict[str, Any],
    observed_base_mz: Any,
    observed_top_ions: Iterable[int],
    family_index: dict[str, dict[str, Any]],
) -> dict[str, Any]:
    row = dict(candidate)
    top5 = list(observed_top_ions)[:5]
    top5_set = set(top5)
    observed_base = _number(observed_base_mz)
    reference_base = _number(row.get("reference_base_peak"))
    reference_diagnostic = [int(ion) for ion in row.get("reference_diagnostic_ions", [])]
    observed_diagnostic = top5_diagnostic_ions(reference_diagnostic, top5)

    if reference_base is None or observed_base is None:
        base_status = "unavailable"
    elif int(reference_base) == int(observed_base):
        base_status = "matched"
    elif int(reference_base) in top5_set and int(observed_base) in reference_diagnostic:
        base_status = "shift_supported"
    else:
        base_status = "conflict"

    base_integer = int(reference_base) if reference_base is not None else None
    additional_expected = [ion for ion in reference_diagnostic if ion != base_integer]
    additional_matched = [ion for ion in additional_expected if ion in top5_set]
    if reference_base is None or not reference_diagnostic:
        compound_status = "unavailable"
    elif base_status == "conflict":
        compound_status = "conflict"
    elif base_status in {"matched", "shift_supported"} and additional_matched:
        compound_status = "supported"
    else:
        compound_status = "not_supported"

    class_code = _text(row.get("reference_class"))
    family = family_index.get(class_code, {})
    family_expected = [int(ion) for ion in family.get("diagnostic_ions", [])]
    family_matched = [ion for ion in family_expected if ion in top5_set]

    molecular_ion = _number(row.get("reference_M_plus"))
    if molecular_ion is None:
        molecular_status = "unavailable"
    elif int(molecular_ion) in top5_set:
        molecular_status = "matched"
    else:
        molecular_status = "not_in_top5"

    family_status, family_base_status, family_reason = _classify_family_ion_evidence(
        row=row,
        class_code=class_code,
        family_name=_text(family.get("family_name")),
        observed_base=observed_base,
        top5=top5,
        family_expected=family_expected,
        family_matched=family_matched,
        molecular_status=molecular_status,
    )
    family_literature_supported = bool(family.get("literature_supported"))
    compound_literature_supported = bool(row.get("literature_supported"))
    literature_support_level = _text(row.get("literature_support_level")) or "Unavailable"
    if (
        not compound_literature_supported
        and family_status == "supported"
        and family_literature_supported
    ):
        literature_support_level = "Family"

    row.update({
        "observed_top5_ions": _join_ions(top5),
        "reference_diagnostic_ions_text": _join_ions(reference_diagnostic),
        "observed_diagnostic_ions": _join_ions(observed_diagnostic),
        "diagnostic_ion_match": diagnostic_ion_match_status(
            reference_diagnostic,
            observed_diagnostic,
        ),
        "base_peak_status": base_status,
        "expected_compound_diagnostic_ions": _join_ions(reference_diagnostic),
        "expected_additional_diagnostic_ions": _join_ions(additional_expected),
        "matched_additional_diagnostic_ions": _join_ions(additional_matched),
        "n_additional_diagnostic_ions_matched": len(additional_matched),
        "compound_ion_status": compound_status,
        "family_name": _text(family.get("family_name")),
        "expected_family_ions": _join_ions(family_expected),
        "matched_family_ions": _join_ions(family_matched),
        "n_family_ions_matched": len(family_matched),
        "family_ion_status": family_status,
        "family_base_peak_status": family_base_status,
        "family_evidence_reason": family_reason,
        "compound_literature_supported": compound_literature_supported,
        "family_literature_supported": family_literature_supported,
        "family_literature_references": _text(family.get("literature_references")),
        "literature_support_level": literature_support_level,
        "M_plus_status": molecular_status,
    })
    return row


def _score(row: dict[str, Any]) -> float:
    return _number(row.get("score_0_1000")) or 0.0


def _ri_order(row: dict[str, Any]) -> tuple[float, float]:
    delta = _number(row.get("abs_delta_RI"))
    return (float("inf") if delta is None else delta, -_score(row))


def _base_compatibility_status(row: dict[str, Any]) -> str:
    reference_status = _text(row.get("base_peak_status"))
    if reference_status in {"matched", "shift_supported"}:
        return reference_status
    observed_base = _number(row.get("observed_base_mz"))
    molecular_weight = _number(row.get("MW"))
    if observed_base is not None and molecular_weight is not None and abs(observed_base - molecular_weight) < 0.5:
        return "molecular_ion_base_peak"
    return ""


def _matched_identification_ions(row: dict[str, Any]) -> list[int]:
    ions: set[int] = set()
    if _text(row.get("compound_ion_status")) == "supported":
        observed_base = _number(row.get("observed_base_mz"))
        if observed_base is not None:
            ions.add(int(observed_base))
        ions.update(parse_reference_ions(row.get("matched_additional_diagnostic_ions")))
    if _text(row.get("M_plus_status")) == "matched":
        molecular_ion = _number(row.get("reference_M_plus") or row.get("MW"))
        if molecular_ion is not None:
            ions.add(int(molecular_ion))
    if _text(row.get("family_ion_status")) == "supported":
        ions.update(parse_reference_ions(row.get("matched_family_ions")))
    return sorted(ions)


def _group_evidence_payload(rows: Iterable[dict[str, Any]]) -> dict[str, Any]:
    items = list(rows)
    ions = sorted({ion for row in items for ion in _matched_identification_ions(row)})
    compound_support = [bool(row.get("literature_supported")) for row in items]
    if items and all(compound_support):
        literature_level = "Compound"
    elif any(compound_support) or any(bool(row.get("family_literature_supported")) for row in items):
        literature_level = "Family"
    else:
        literature_level = "Unavailable"
    ri_supported = bool(items) and all(_text(row.get("RI_status")) == "supported" for row in items)
    compound_ions = any(_text(row.get("compound_ion_status")) == "supported" for row in items)
    family_ions = any(_text(row.get("family_ion_status")) == "supported" for row in items)
    evidence_parts = ["Spectral"]
    if ri_supported:
        evidence_parts.append("RI")
    if compound_ions:
        evidence_parts.append("Compound Ions")
    elif family_ions:
        evidence_parts.append("Family Ions")
    evidence = " + ".join(evidence_parts) + " (Isomer Unresolved)"
    return {
        "selected_literature_support_level": literature_level,
        "selected_matched_diagnostic_ions": _join_ions(ions),
        "selected_identification_evidence": evidence,
    }


def _has_complete_compound_evidence(row: dict[str, Any]) -> bool:
    """Return whether a candidate has orthogonal compound-level support."""
    if not bool(row.get("literature_supported")):
        return False
    return (
        _text(row.get("compound_ion_status")) == "supported"
        or bool(_base_compatibility_status(row))
    )


def _candidate_identity(row: dict[str, Any]) -> str:
    inchikey = _text(row.get("InChIKey")).upper()
    if inchikey:
        return f"inchikey:{inchikey}"
    cas = _normalized_cas(row.get("CAS") or row.get("reference_CAS"))
    if cas:
        return f"cas:{cas}"
    return f"name:{_normalized_name(row.get('Name'))}|formula:{_text(row.get('Formula')).upper()}"


def _candidate_formula(row: dict[str, Any]) -> str:
    return _text(
        row.get("resolved_formula")
        or row.get("Formula")
        or row.get("reference_formula")
    ).upper()


def _shared_candidate_formula(rows: Iterable[dict[str, Any]]) -> str:
    items = list(rows)
    formulas = {_candidate_formula(row) for row in items}
    if len(items) < 2 or "" in formulas or len(formulas) != 1:
        return ""
    return next(iter(formulas))


def _candidate_group_kind(rows: Iterable[dict[str, Any]]) -> str:
    items = _unique_candidates(rows)
    if len(items) < 2:
        return "single_identity"
    return "isomer_group" if _shared_candidate_formula(items) else "ambiguous_candidates"


def _has_specific_ion_support(row: dict[str, Any]) -> bool:
    return (
        _text(row.get("M_plus_status")) == "matched"
        or _text(row.get("compound_ion_status")) == "supported"
    )


def _candidate_delta(row: dict[str, Any]) -> float | None:
    delta = _number(row.get("delta_RI"))
    if delta is not None:
        return abs(delta)
    return _number(row.get("abs_delta_RI"))


def _meets_best_candidate_rule(row: dict[str, Any]) -> bool:
    delta = _candidate_delta(row)
    return (
        _score(row) >= 950
        and _text(row.get("RI_status")) == "supported"
        and delta is not None
        and delta <= 20
        and _has_specific_ion_support(row)
        and _text(row.get("base_peak_status")) != "conflict"
        and _text(row.get("compound_ion_status")) != "conflict"
        and _text(row.get("M_plus_status")) != "conflict"
    )


def _automatic_best_candidate(rows: Iterable[dict[str, Any]]) -> dict[str, Any] | None:
    qualified = [row for row in _unique_candidates(rows) if _meets_best_candidate_rule(row)]
    return qualified[0] if len(qualified) == 1 else None


def _unresolved_group_decision(
    base: dict[str, Any],
    rows: Iterable[dict[str, Any]],
) -> dict[str, Any]:
    items = _unique_candidates(rows)
    winner = _automatic_best_candidate(items)
    if winner is not None:
        return {
            **base,
            "automated_status": (
                "putative_best_candidate_high_confidence"
                if bool(winner.get("literature_supported"))
                else "putative_best_candidate_medium_confidence"
            ),
            "automatic_identification_eligible": True,
            "selected_candidate": winner.get("Name", ""),
            "selected_formula": _candidate_formula(winner),
            "selected_CAS": winner.get("CAS", ""),
            "selected_library": winner.get("library", ""),
            "selected_match_score": winner.get("score_0_1000", ""),
            "selected_reference_RI": winner.get("reference_RI", ""),
            "selected_delta_RI": winner.get("delta_RI", ""),
            "selected_base_peak_status": _base_compatibility_status(winner)
            or winner.get("base_peak_status", ""),
            "selected_compound_ion_status": winner.get("compound_ion_status", ""),
            "selected_matched_compound_ions": winner.get(
                "matched_additional_diagnostic_ions", ""
            ),
            "selected_family": winner.get("family_name", ""),
            "selected_family_ion_status": winner.get("family_ion_status", ""),
            "selected_matched_family_ions": winner.get("matched_family_ions", ""),
            "selected_M_plus_status": winner.get("M_plus_status", ""),
            "selected_literature_support_level": winner.get(
                "literature_support_level", "Unavailable"
            ),
            "selected_matched_diagnostic_ions": _join_ions(
                _matched_identification_ions(winner)
            ),
            "selected_identification_evidence": (
                "Spectral + RI + Compound Ions"
                if _text(winner.get("compound_ion_status")) == "supported"
                else "Spectral + RI + Molecular Ion"
            ),
            "candidate_group": " | ".join(
                dict.fromkeys(
                    _text(row.get("Name"))
                    for row in items
                    if _text(row.get("Name"))
                )
            ),
        }

    names = [
        _text(row.get("Name"))
        for row in items
        if _text(row.get("Name"))
    ]
    group_kind = _candidate_group_kind(items)
    payload = _group_evidence_payload(items)
    if group_kind == "isomer_group":
        if not payload["selected_matched_diagnostic_ions"]:
            observed_ions = sorted({
                ion
                for row in items
                for ion in parse_observed_ions(row.get("observed_top5_ions"))
            })
            payload["selected_matched_diagnostic_ions"] = _join_ions(observed_ions)
        return {
            **base,
            "automated_status": "isomer_group",
            "automatic_identification_eligible": True,
            "selected_formula": _shared_candidate_formula(items),
            "candidate_group": " | ".join(dict.fromkeys(names)),
            **payload,
        }
    return {
        **base,
        "automated_status": "ambiguous_candidates_manual_review",
        "automatic_identification_eligible": False,
        "candidate_group": " | ".join(dict.fromkeys(names)),
        **payload,
    }


def _candidates_equivalent(left: dict[str, Any], right: dict[str, Any]) -> bool:
    left_key = _text(left.get("InChIKey")).upper()
    right_key = _text(right.get("InChIKey")).upper()
    if left_key and right_key and left_key == right_key:
        return True
    left_cas = _normalized_cas(left.get("CAS") or left.get("reference_CAS"))
    right_cas = _normalized_cas(right.get("CAS") or right.get("reference_CAS"))
    if left_cas and right_cas and left_cas == right_cas:
        return True
    left_formula = _text(left.get("Formula") or left.get("reference_formula")).upper()
    right_formula = _text(right.get("Formula") or right.get("reference_formula")).upper()
    if not left_formula or not right_formula or left_formula != right_formula:
        return False
    names_left = {_normalized_name(left.get("Name")), _normalized_name(left.get("reference_compound"))} - {""}
    names_right = {_normalized_name(right.get("Name")), _normalized_name(right.get("reference_compound"))} - {""}
    return bool(names_left & names_right)


def _unique_candidates(rows: Iterable[dict[str, Any]]) -> list[dict[str, Any]]:
    items = list(rows)
    parents = list(range(len(items)))

    def find(index: int) -> int:
        while parents[index] != index:
            parents[index] = parents[parents[index]]
            index = parents[index]
        return index

    def union(left: int, right: int) -> None:
        left_root, right_root = find(left), find(right)
        if left_root != right_root:
            parents[right_root] = left_root

    for left in range(len(items)):
        for right in range(left + 1, len(items)):
            if _candidates_equivalent(items[left], items[right]):
                union(left, right)
    groups: dict[int, list[dict[str, Any]]] = {}
    for index, row in enumerate(items):
        groups.setdefault(find(index), []).append(row)

    def evidence_rank(row: dict[str, Any]) -> tuple[int, int, int, int, float]:
        return (
            int(_text(row.get("compound_ion_status")) == "supported"),
            int(_text(row.get("RI_status")) == "supported"),
            int(_text(row.get("family_ion_status")) == "supported"),
            int(bool(row.get("literature_supported"))),
            _score(row),
        )

    evidence_fields = {
        "RI_status", "reference_RI", "delta_RI", "abs_delta_RI",
        "literature_supported", "compound_literature_supported", "family_literature_supported",
        "family_literature_references", "literature_support_level",
        "reference_compound", "reference_formula",
        "reference_CAS", "reference_class", "reference_base_peak",
        "reference_diagnostic_ions", "reference_M_plus", "base_peak_status",
        "expected_compound_diagnostic_ions", "expected_additional_diagnostic_ions",
        "matched_additional_diagnostic_ions", "n_additional_diagnostic_ions_matched",
        "compound_ion_status", "family_name", "expected_family_ions",
        "matched_family_ions", "n_family_ions_matched", "family_ion_status",
        "family_base_peak_status", "family_evidence_reason",
        "M_plus_status",
    }
    merged_groups: list[dict[str, Any]] = []
    for group in groups.values():
        highest_score = max(group, key=_score)
        evidence_source = max(group, key=evidence_rank)
        merged = dict(highest_score)
        for field in evidence_fields:
            if field in evidence_source:
                merged[field] = evidence_source[field]
        if _text(evidence_source.get("reference_compound")):
            merged["Name"] = evidence_source["reference_compound"]
        elif _text(evidence_source.get("Name")):
            merged["Name"] = evidence_source["Name"]
        if not _text(merged.get("Formula")) and _text(evidence_source.get("reference_formula")):
            merged["Formula"] = evidence_source["reference_formula"]
        if not _text(merged.get("CAS")) and _text(evidence_source.get("reference_CAS")):
            merged["CAS"] = evidence_source["reference_CAS"]
        merged_groups.append(merged)
    return merged_groups


def select_feature_decision(
    audited: Iterable[dict[str, Any]],
    ri_calculation_status: str = "interpolated",
    spectral_auto_threshold: float = 850.0,
    decision_policy: str = "legacy_ion_first",
) -> dict[str, Any]:
    decision_policies = {"legacy_ion_first", "ri_first_conservative"}
    if decision_policy not in decision_policies:
        raise ValueError(
            f"Unknown decision_policy {decision_policy!r}; "
            f"expected one of {sorted(decision_policies)}"
        )
    rows = list(audited)
    unique_rows = _unique_candidates(rows)
    supported = sorted((row for row in unique_rows if row.get("RI_status") == "supported"), key=_ri_order)
    weak = sorted((row for row in unique_rows if row.get("RI_status") == "weak_support"), key=_ri_order)
    comparable = [row for row in unique_rows if row.get("RI_status") != "unavailable"]
    unavailable = sorted((row for row in unique_rows if row.get("RI_status") == "unavailable"), key=lambda row: -_score(row))
    eligible = [row for row in unique_rows if _score(row) >= spectral_auto_threshold]
    eligible_ids = {_candidate_identity(row) for row in eligible}
    supported_eligible = [row for row in supported if _candidate_identity(row) in eligible_ids]
    unavailable_eligible = [row for row in unavailable if _candidate_identity(row) in eligible_ids]

    base = {
        "n_candidates_retained": len(rows),
        "n_unique_candidates": len(unique_rows),
        "n_RI_supported_candidates": len(supported),
        "n_auto_eligible_candidates": len(eligible),
        "spectral_auto_threshold": spectral_auto_threshold,
        "automatic_identification_eligible": False,
        "RI_ambiguity_gap": "",
        "selected_candidate": "",
        "selected_formula": "",
        "selected_CAS": "",
        "selected_library": "",
        "selected_match_score": "",
        "selected_reference_RI": "",
        "selected_delta_RI": "",
        "selected_base_peak_status": "",
        "selected_compound_ion_status": "",
        "selected_matched_compound_ions": "",
        "selected_family": "",
        "selected_family_ion_status": "",
        "selected_matched_family_ions": "",
        "selected_M_plus_status": "",
        "selected_literature_support_level": "Unavailable",
        "selected_matched_diagnostic_ions": "",
        "selected_identification_evidence": "",
    }
    winner: dict[str, Any]
    candidate_group = ""
    ambiguity_value: float | str = ""
    selected_base_peak_status = ""
    compound_supported = [
        row for row in eligible
        if _text(row.get("compound_ion_status")) == "supported"
        and _text(row.get("base_peak_status")) != "conflict"
    ]
    if ri_calculation_status != "interpolated":
        if len(compound_supported) > 1:
            return _unresolved_group_decision(base, compound_supported)
        if len(compound_supported) == 1:
            winner = compound_supported[0]
            selected_base_peak_status = _base_compatibility_status(winner)
            status = "putative_compound_ion_supported"
        else:
            return {**base, "automated_status": "manual_review_RI_unavailable"}
    elif len(supported_eligible) > 1:
        return _unresolved_group_decision(base, supported_eligible)
    elif decision_policy == "ri_first_conservative" and len(supported_eligible) == 1:
        winner = supported_eligible[0]
        selected_base_peak_status = _base_compatibility_status(winner)
        if (
            _text(winner.get("compound_ion_status")) == "conflict"
            or _text(winner.get("base_peak_status")) == "conflict"
        ):
            status = "RI_supported_manual_review"
        elif _text(winner.get("compound_ion_status")) == "supported":
            status = "putative_compound_multiple_evidence"
        else:
            status = "putative_compound_RI_selected"
    elif len(compound_supported) > 1:
        compound_ri_supported = [
            row for row in compound_supported if row.get("RI_status") == "supported"
        ]
        if len(compound_ri_supported) == 1:
            winner = compound_ri_supported[0]
            selected_base_peak_status = _base_compatibility_status(winner)
            status = "putative_compound_multiple_evidence"
        else:
            ambiguous = compound_ri_supported if compound_ri_supported else compound_supported
            return _unresolved_group_decision(base, ambiguous)
    elif len(compound_supported) == 1:
        winner = compound_supported[0]
        selected_base_peak_status = _base_compatibility_status(winner)
        winner_ri_status = _text(winner.get("RI_status"))
        if winner_ri_status == "weak_support":
            status = "manual_review_weak_RI_support"
        elif winner_ri_status == "mismatch":
            status = "RI_conflict_manual_review"
        elif winner_ri_status == "supported":
            status = "putative_compound_multiple_evidence"
        else:
            status = "putative_compound_ion_supported"
    elif supported_eligible:
        winner = supported_eligible[0]
        selected_base_peak_status = _base_compatibility_status(winner)
        if _text(winner.get("compound_ion_status")) == "conflict" or _text(winner.get("base_peak_status")) == "conflict":
            status = "RI_supported_manual_review"
        else:
            status = "putative_compound_RI_selected"
    elif supported:
        winner = supported[0]
        status = "manual_review_spectral_below_threshold"
    else:
        family_supported = [
            row for row in eligible
            if _text(row.get("family_ion_status")) == "supported"
            and _text(row.get("family_name"))
        ]
        family_names = list(dict.fromkeys(
            _text(row.get("family_name")) for row in family_supported
            if _text(row.get("family_name"))
        ))
        if family_supported and len(family_names) == 1:
            winner = sorted(family_supported, key=lambda row: -_score(row))[0]
            status = "putative_family_ion_supported"
            candidate_group = " | ".join(dict.fromkeys(
                _text(row.get("Name")) for row in family_supported if _text(row.get("Name"))
            ))
        elif family_supported:
            return _unresolved_group_decision(base, family_supported)
        elif weak:
            weak_eligible = [row for row in weak if _candidate_identity(row) in eligible_ids]
            winner = weak_eligible[0] if weak_eligible else weak[0]
            status = "manual_review_weak_RI_support"
        elif comparable:
            winner = sorted(comparable, key=_ri_order)[0]
            status = "RI_conflict_manual_review"
        else:
            if unavailable_eligible:
                winner = unavailable_eligible[0]
                status = "spectral_literature_manual_review" if winner.get("literature_supported") else "manual_review_insufficient_evidence"
            else:
                winner = unavailable[0] if unavailable else {}
                status = "manual_review_spectral_below_threshold"

    selected_name = winner.get("Name", "")
    if status == "putative_family_ion_supported":
        selected_name = winner.get("family_name", "")

    return {
        **base,
        "automated_status": status,
        "automatic_identification_eligible": _score(winner) >= spectral_auto_threshold,
        "RI_ambiguity_gap": ambiguity_value,
        "selected_candidate": selected_name,
        "selected_formula": winner.get("Formula", ""),
        "selected_CAS": winner.get("CAS", ""),
        "selected_library": winner.get("library", ""),
        "selected_match_score": winner.get("score_0_1000", ""),
        "selected_reference_RI": winner.get("reference_RI", ""),
        "selected_delta_RI": winner.get("delta_RI", ""),
        "selected_base_peak_status": selected_base_peak_status or winner.get("base_peak_status", ""),
        "selected_compound_ion_status": winner.get("compound_ion_status", ""),
        "selected_matched_compound_ions": winner.get("matched_additional_diagnostic_ions", ""),
        "selected_family": winner.get("family_name", ""),
        "selected_family_ion_status": winner.get("family_ion_status", ""),
        "selected_matched_family_ions": winner.get("matched_family_ions", ""),
        "selected_M_plus_status": winner.get("M_plus_status", ""),
        "selected_literature_support_level": (
            "Family" if status == "putative_family_ion_supported"
            else winner.get("literature_support_level", "Unavailable")
        ),
        "selected_matched_diagnostic_ions": _join_ions(_matched_identification_ions(winner)),
        "selected_reference_diagnostic_ions": winner.get(
            "reference_diagnostic_ions_text", ""
        ),
        "selected_observed_diagnostic_ions": winner.get(
            "observed_diagnostic_ions", ""
        ),
        "selected_diagnostic_ion_match": winner.get(
            "diagnostic_ion_match", "Not defined"
        ),
        "selected_identification_evidence": (
            "Spectral + Compound Ions"
            if status == "putative_compound_ion_supported"
            else ""
        ),
        "candidate_group": candidate_group,
    }


def _candidate_rank(row: dict[str, Any]) -> int:
    rank = _number(row.get("rank"))
    return int(rank) if rank is not None and rank >= 1 else 10**9


def _decision_is_auto_reportable(decision: dict[str, Any]) -> bool:
    status = _text(decision.get("automated_status"))
    selected = _text(decision.get("selected_candidate"))
    group = _text(decision.get("candidate_group"))
    if status in {"isomer_group", "isomer_group_or_manual_review"}:
        return bool(selected or group)
    return bool(
        selected
        and (
            status.startswith("putative_")
            or status.startswith("homologous_series_")
        )
    )


def select_two_stage_feature_decision(
    audited: Iterable[dict[str, Any]],
    ri_calculation_status: str = "interpolated",
    spectral_auto_threshold: float = 850.0,
    primary_top_n: int = 5,
    rescue_top_n: int = 20,
) -> dict[str, Any]:
    if primary_top_n < 1:
        raise ValueError("primary_top_n must be >= 1")
    if rescue_top_n < primary_top_n:
        raise ValueError("rescue_top_n must be >= primary_top_n")

    rows = list(audited)
    primary_rows = [row for row in rows if _candidate_rank(row) <= primary_top_n]
    rescue_rows = [row for row in rows if _candidate_rank(row) <= rescue_top_n]
    primary = select_feature_decision(
        primary_rows,
        ri_calculation_status=ri_calculation_status,
        spectral_auto_threshold=spectral_auto_threshold,
        decision_policy="legacy_ion_first",
    )
    primary_audit = {
        "primary_automated_status": primary.get("automated_status", ""),
        "primary_selected_candidate": primary.get("selected_candidate", ""),
    }
    if _decision_is_auto_reportable(primary):
        return {
            **primary,
            **primary_audit,
            "decision_round": "top5_primary",
            "rescue_attempted": False,
            "rescue_automated_status": "",
            "rescue_selected_candidate": "",
        }

    conservative_rescue_rows: list[dict[str, Any]] = []
    for row in rescue_rows:
        if (
            _text(row.get("RI_status")) == "unavailable"
            and _text(row.get("compound_ion_status")) == "supported"
            and not _base_compatibility_status(row)
        ):
            row = {**row, "compound_ion_status": "unavailable"}
        conservative_rescue_rows.append(row)
    rescue = select_feature_decision(
        conservative_rescue_rows,
        ri_calculation_status=ri_calculation_status,
        spectral_auto_threshold=spectral_auto_threshold,
        decision_policy="ri_first_conservative",
    )
    return {
        **rescue,
        **primary_audit,
        "decision_round": (
            "top20_rescue"
            if _decision_is_auto_reportable(rescue)
            else "manual_after_top20"
        ),
        "rescue_attempted": True,
        "rescue_automated_status": rescue.get("automated_status", ""),
        "rescue_selected_candidate": rescue.get("selected_candidate", ""),
    }


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="RI-guided adjudication of all retained NIST/MoNA Top-N candidates")
    parser.add_argument("--top-hits", required=True, type=Path)
    parser.add_argument("--feature-metadata", required=True, type=Path)
    parser.add_argument("--alkanes", required=True, type=Path)
    parser.add_argument("--reference-workbook", required=True, type=Path)
    parser.add_argument("--candidate-output", required=True, type=Path)
    parser.add_argument("--decision-output", required=True, type=Path)
    parser.add_argument("--cache-dir", required=True, type=Path)
    parser.add_argument("--pubchem-cache-dir", type=Path)
    parser.add_argument("--support-window", type=float, default=20.0)
    parser.add_argument("--weak-window", type=float, default=50.0)
    parser.add_argument("--spectral-auto-threshold", type=float, default=850.0)
    parser.add_argument(
        "--decision-policy",
        choices=(
            "legacy_ion_first",
            "ri_first_conservative",
            "two_stage_top20_rescue",
        ),
        default="legacy_ion_first",
    )
    parser.add_argument("--primary-top-n", type=int, default=5)
    parser.add_argument("--rescue-top-n", type=int, default=20)
    parser.add_argument("--timeout", type=float, default=20.0)
    parser.add_argument("--delay", type=float, default=0.12)
    parser.add_argument("--pubchem-timeout", type=float, default=20.0)
    parser.add_argument("--pubchem-delay", type=float, default=0.12)
    parser.add_argument("--offline", action="store_true")
    parser.add_argument("--pubchem-offline", action="store_true")
    return parser.parse_args(argv)


def main() -> None:
    args = parse_args()
    candidates, decisions = run_identification(
        top_hits_path=args.top_hits,
        feature_metadata_path=args.feature_metadata,
        alkane_path=args.alkanes,
        reference_workbook_path=args.reference_workbook,
        candidate_output_path=args.candidate_output,
        decision_output_path=args.decision_output,
        cache_dir=args.cache_dir,
        pubchem_cache_dir=args.pubchem_cache_dir,
        online=not args.offline,
        pubchem_online=not args.pubchem_offline,
        timeout=args.timeout,
        delay=args.delay,
        pubchem_timeout=args.pubchem_timeout,
        pubchem_delay=args.pubchem_delay,
        support_window=args.support_window,
        weak_window=args.weak_window,
        spectral_auto_threshold=args.spectral_auto_threshold,
        decision_policy=args.decision_policy,
        primary_top_n=args.primary_top_n,
        rescue_top_n=args.rescue_top_n,
    )
    print(f"Wrote {args.candidate_output} ({len(candidates)} retained candidate rows)")
    print(f"Wrote {args.decision_output} ({len(decisions)} feature decisions)")
main.select_feature_decision = select_feature_decision
main.select_two_stage_feature_decision = select_two_stage_feature_decision
main.load_reference_workbook = load_reference_workbook
main.resolve_feature_ri = resolve_feature_ri
main.detect_homologous_series = detect_homologous_series
main.parse_top_ions = parse_top_ions
main.parse_reference_ions = parse_reference_ions
main.top5_diagnostic_ions = top5_diagnostic_ions
main.diagnostic_ion_match_status = diagnostic_ion_match_status
main.attach_ion_evidence = attach_ion_evidence
main.run_identification = run_identification
main.parse_args = parse_args
